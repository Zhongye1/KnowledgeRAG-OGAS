"""Office 文档直读处理器（docx/pptx → Markdown，ragf-design D13 首发子集）。

docx/pptx 走“文本直读”而非 OCR：按文档元素顺序抽取段落/标题/表格/形状文本，
产出 Markdown 兼容文本。布局级版面解析（docling 路线）作为后续引擎选项。
"""

from __future__ import annotations

import io

from typing import Any, ClassVar

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.src.app.ingest.parser.base import BaseDocumentProcessor, DocumentParseError
from backend.src.app.ingest.parser.registry import register

_HEADING_PREFIX = 'Heading '


def _docx_to_markdown(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith('}p'):
            _append_docx_paragraph(Paragraph(child, document), lines)
        elif child.tag.endswith('}tbl'):
            _append_docx_table(Table(child, document), lines)
    return '\n'.join(lines).strip()


def _append_docx_paragraph(paragraph: Paragraph, lines: list[str]) -> None:
    text = ''.join(run.text for run in paragraph.runs).strip()
    if not text:
        return
    style_name = (paragraph.style.name if paragraph.style is not None else '') or ''
    if style_name.startswith(_HEADING_PREFIX):
        level = style_name[len(_HEADING_PREFIX) :].strip()
        level = int(level) if level.isdigit() else 1
        lines.append(f'{"#" * min(max(level, 1), 6)} {text}')
    else:
        lines.append(text)


def _append_docx_table(table: Table, lines: list[str]) -> None:
    rendered = []
    for row in table.rows:
        cells = [' '.join(cell.text.split()) for cell in row.cells]
        rendered.append('| ' + ' | '.join(cells) + ' |')
    if rendered:
        lines.append('\n'.join(rendered))


def _pptx_to_markdown(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    lines: list[str] = []
    for slide in presentation.slides:
        slide_lines: list[str] = []
        for shape in slide.shapes:
            _collect_shape_text(shape, slide_lines, MSO_SHAPE_TYPE)
        if slide_lines:
            lines.append('\n'.join(slide_lines))
    return '\n\n'.join(lines).strip()


def _collect_shape_text(shape: Any, lines: list[str], shape_types: Any) -> None:
    shape_type = getattr(shape, 'shape_type', None)
    if shape_type == shape_types.GROUP:
        for child in shape.shapes:
            _collect_shape_text(child, lines, shape_types)
        return
    if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
        lines.append(shape.text_frame.text.strip())
    elif getattr(shape, 'has_table', False):
        rows = []
        for row in shape.table.rows:
            cells = [' '.join(cell.text.split()) for cell in row.cells]
            rows.append('| ' + ' | '.join(cells) + ' |')
        lines.append('\n'.join(rows))


@register
class OfficeTextProcessor(BaseDocumentProcessor):
    """docx/pptx 文本直读（Markdown 化，无 OCR 依赖）。"""

    service_name: ClassVar[str] = 'office_text'
    display_name: ClassVar[str] = 'Office 文本直读'
    supported_extensions: ClassVar[list[str]] = ['.docx', '.pptx']

    async def parse_bytes(self, data: bytes, filename: str, params: dict[str, Any] | None = None) -> str:
        if not data:
            raise DocumentParseError('文件内容为空', service_name=self.service_name, error_code='empty_content')
        try:
            markdown = _docx_to_markdown(data) if filename.lower().endswith('.docx') else _pptx_to_markdown(data)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError(
                f'Office 文档解析失败: {exc}',
                service_name=self.service_name,
                error_code='parse_error',
            ) from exc
        if not markdown:
            raise DocumentParseError(
                'Office 文档无可抽取文本', service_name=self.service_name, error_code='empty_content'
            )
        return markdown

    async def check_health(self) -> dict[str, Any]:
        return {'status': 'healthy', 'message': 'Office 文本直读无需外部服务', 'details': {}}
