"""解析层单元测试（Yuxi knowledge/parser 移植，ragf-design §11）。

覆盖：direct text（md/txt/csv）解码路由、不支持格式/未注册引擎的明确错误、注册表内容。
外部引擎（MinerU）不做网络依赖断言，仅验证实例化与注册。
"""

from __future__ import annotations

import asyncio

from typing import TYPE_CHECKING

import pytest

from backend.src.app.ingest.parser import DocumentProcessorFactory, parse_document
from backend.src.app.ingest.parser.base import DocumentParseError, ProcessorUnavailableError

if TYPE_CHECKING:
    from collections.abc import Coroutine
    from typing import Any


def _run[T](awaitable: Coroutine[Any, Any, T]) -> T:
    return asyncio.run(awaitable)


def test_parse_markdown_passthrough() -> None:
    """md 直读：原文透传。"""
    text = '# 标题\n\n正文段落。\n'
    assert _run(parse_document(text.encode(), 'demo.md', {})) == text.strip()


def test_parse_txt_gb18030_fallback() -> None:
    """txt：UTF-8 失败时按 GB18030 解码。"""
    text = '中文内容行一\n行二'
    assert _run(parse_document(text.encode('gb18030'), 'demo.txt', {})) == text


def test_parse_csv_passthrough() -> None:
    """csv：原文透传（QA/结构化解析在分块层处理）。"""
    raw = 'a,b\n1,2\n'
    assert _run(parse_document(raw.encode(), 'demo.csv', {})) == raw.strip()


def test_parse_unsupported_format_raises() -> None:
    """xlsx 等未接入引擎格式：明确 unavailable 错误（D13 子集外拒绝）。"""
    with pytest.raises(ProcessorUnavailableError):
        _run(parse_document(b'x', 'demo.xlsx', {}))


def test_parse_docx_office_text() -> None:
    """docx 直读：标题转 Markdown 井号、表格转管线行。"""
    import io

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading('一期标题', level=1)
    document.add_paragraph('正文段落。')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '甲'
    table.cell(0, 1).text = '乙'
    buffer = io.BytesIO()
    document.save(buffer)
    markdown = _run(parse_document(buffer.getvalue(), 'demo.docx', {}))
    assert '# 一期标题' in markdown
    assert '正文段落。' in markdown
    assert '| 甲 | 乙 |' in markdown


def test_parse_pptx_office_text() -> None:
    """pptx 直读：形状文本按 slide 顺序抽取。"""
    import io

    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches

    presentation = PptxPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text = '幻灯片要点'
    buffer = io.BytesIO()
    presentation.save(buffer)
    markdown = _run(parse_document(buffer.getvalue(), 'demo.pptx', {}))
    assert '幻灯片要点' in markdown


def _cjk_font_path() -> str | None:
    """查找可渲染中文的系统字体（缺失则跳过真实 OCR 用例）。"""
    import glob

    for pattern in (
        '/usr/share/fonts/**/wqy-zenhei.ttc',
        '/usr/share/fonts/**/*CJK*.ttc',
        '/usr/share/fonts/**/*CJK*.otf',
        '/usr/share/fonts/**/Noto*.ttc',
    ):
        hits = glob.glob(pattern, recursive=True)
        if hits:
            return hits[0]
    return None


def test_parse_png_via_rapid_ocr() -> None:
    """rapid_ocr 真实 OCR：中文 PNG 应识别出文本（无 CJK 字体/依赖则跳过）。"""
    import io

    from PIL import Image, ImageDraw, ImageFont

    font_path = _cjk_font_path()
    if font_path is None:
        pytest.skip('无中文字体，跳过真实 OCR 用例')
    try:
        import rapidocr  # ruff: ignore[unused-import]
    except ImportError:
        pytest.skip('rapidocr 未安装')

    image = Image.new('RGB', (420, 80), 'white')
    ImageDraw.Draw(image).text((10, 20), 'RAG混合检索测试', font=ImageFont.truetype(font_path, 32), fill='black')
    buffer = io.BytesIO()
    image.save(buffer, format='PNG')
    markdown = _run(parse_document(buffer.getvalue(), 'demo.png', {'ocr_engine': 'rapid_ocr'}))
    assert '检索' in markdown and 'RAG' in markdown


def test_parse_document_fallback_to_rapid_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    """OCR 主引擎失败 → 自动降级 rapid_ocr（M4 兜底，记录实际引擎）。"""
    from backend.src.app.ingest.parser import DocumentProcessorFactory
    from backend.src.app.ingest.parser.registry import PROCESSORS

    class FailingMineru:
        service_name = 'mineru'

        async def parse_bytes(self, data: bytes, filename: str, params: dict | None = None) -> str:
            raise DocumentParseError('容器不可达', service_name='mineru', error_code='parse_error')

        async def check_health(self) -> dict:
            return {}

    class WorkingRapid:
        service_name = 'rapid_ocr'

        async def parse_bytes(self, data: bytes, filename: str, params: dict | None = None) -> str:
            return '兜底内容'

        async def check_health(self) -> dict:
            return {}

    def fake_get(engine_id: str, **kwargs: object) -> object:
        if engine_id == 'mineru':
            return FailingMineru()
        if engine_id == 'rapid_ocr':
            return WorkingRapid()
        raise ProcessorUnavailableError(f'未注册: {engine_id}')

    monkeypatch.setattr(DocumentProcessorFactory, 'get_processor', staticmethod(fake_get))  # type: ignore[method-assign]
    assert 'mineru' in PROCESSORS and 'rapid_ocr' in PROCESSORS
    markdown, engine = _run(
        DocumentProcessorFactory.parse_document_with_fallback(b'pdf-bytes', 'demo.pdf', {'ocr_engine': 'mineru'})
    )
    assert markdown == '兜底内容'
    assert engine == 'rapid_ocr'


def test_parse_unknown_engine_raises() -> None:
    """settings/params 指定未注册 OCR 引擎：明确错误。"""
    with pytest.raises(ProcessorUnavailableError):
        _run(DocumentProcessorFactory.parse_document(b'x', 'demo.pdf', {'ocr_engine': 'not-registered'}))


def test_registry_engines() -> None:
    """注册表含首发引擎（direct_text 常驻；mineru 为 OCR 默认）。"""
    engines = DocumentProcessorFactory.get_available_processors()
    assert 'direct_text' in engines
    assert 'office_text' in engines
    assert 'mineru' in engines
    assert 'rapid_ocr' in engines
